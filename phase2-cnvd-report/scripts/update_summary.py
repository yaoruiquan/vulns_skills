#!/usr/bin/env python3
"""
更新漏洞汇总表

将提交的 CNVD/CNNVD 漏洞信息添加到同一张 xlsx 汇总表中。
"""

import argparse
import os
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook


SKILL_ROOT = Path(__file__).resolve().parents[1]
ENV_FILE = SKILL_ROOT / ".env"
DEFAULT_SUMMARY_PATH = "/Users/yao/Documents/网安- AI应用开发/监管上报/汇总表/漏洞汇总表.xlsx"


def load_env() -> None:
    """加载 skill 根目录下的 .env，不覆盖已存在环境变量。"""
    if not ENV_FILE.exists():
        return
    for raw_line in ENV_FILE.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        value = os.path.expandvars(value.strip().strip('"').strip("'"))
        if key and key not in os.environ:
            os.environ[key] = value


def get_summary_path() -> str:
    load_env()
    return os.environ.get("SUMMARY_TABLE_PATH", DEFAULT_SUMMARY_PATH)


def ensure_dir(filepath: str) -> None:
    dir_path = os.path.dirname(filepath)
    if dir_path and not os.path.exists(dir_path):
        os.makedirs(dir_path, exist_ok=True)


def create_summary_file(filepath: str):
    ensure_dir(filepath)
    wb = Workbook()
    ws = wb.active
    ws.title = "漏洞汇总"

    headers = ["漏洞标题", "影响厂商", "漏洞编号", "提交人员", "上报CNVD编号", "上报CNNVD编号", "上报日期"]
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)

    column_widths = [60, 20, 15, 15, 20, 20, 15]
    for col, width in enumerate(column_widths, 1):
        ws.column_dimensions[chr(64 + col)].width = width

    wb.save(filepath)
    print(f"创建汇总表: {filepath}")
    return wb


def find_row_by_das_id(ws, das_id: str):
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=3).value == das_id:
            return row
    return None


def update_summary(
    title: str,
    vendor: str,
    das_id: str,
    submitter: str,
    cnvd_id: str,
    cnnvd_id: str,
    date: str,
    filepath: str = None,
) -> None:
    if filepath is None:
        filepath = get_summary_path()

    if not os.path.exists(filepath):
        wb = create_summary_file(filepath)
        ws = wb.active
    else:
        wb = load_workbook(filepath)
        ws = wb.active

    row = find_row_by_das_id(ws, das_id)
    if row:
        print(f"更新现有记录 (行 {row}): {das_id}")
    else:
        row = ws.max_row + 1
        print(f"添加新记录 (行 {row}): {das_id}")

    existing_cnvd_id = ws.cell(row=row, column=5).value or ""
    existing_cnnvd_id = ws.cell(row=row, column=6).value or ""

    ws.cell(row=row, column=1, value=title)
    ws.cell(row=row, column=2, value=vendor)
    ws.cell(row=row, column=3, value=das_id)
    ws.cell(row=row, column=4, value=submitter)
    final_cnvd_id = cnvd_id if cnvd_id else existing_cnvd_id
    final_cnnvd_id = cnnvd_id if cnnvd_id else existing_cnnvd_id
    ws.cell(row=row, column=5, value=final_cnvd_id)
    ws.cell(row=row, column=6, value=final_cnnvd_id)
    ws.cell(row=row, column=7, value=date)

    wb.save(filepath)
    print(f"汇总表已更新: {filepath}")
    print(f"  漏洞标题: {title}")
    print(f"  影响厂商: {vendor}")
    print(f"  漏洞编号: {das_id}")
    print(f"  提交人员: {submitter}")
    print(f"  CNVD编号: {final_cnvd_id or '(空)'}")
    print(f"  CNNVD编号: {final_cnnvd_id or '(空)'}")
    print(f"  上报日期: {date}")


def main() -> int:
    parser = argparse.ArgumentParser(description="更新漏洞汇总表")
    parser.add_argument("--title", required=True, help="漏洞标题")
    parser.add_argument("--vendor", default="", help="影响厂商")
    parser.add_argument("--das-id", required=True, help="漏洞编号 (DAS-ID)")
    parser.add_argument("--submitter", default="", help="提交人员")
    parser.add_argument("--cnvd-id", default="", help="上报 CNVD 编号")
    parser.add_argument("--cnnvd-id", default="", help="上报 CNNVD 编号")
    parser.add_argument("--platform", choices=["CNVD", "CNNVD"], default="", help="平台名；配合 --platform-id 自动写入对应列")
    parser.add_argument("--platform-id", default="", help="平台编号；配合 --platform 自动写入 CNVD/CNNVD 对应列")
    parser.add_argument("--date", default=None, help="上报日期 (默认今天)")
    parser.add_argument("--filepath", default=None, help="汇总表文件路径；默认读取 SUMMARY_TABLE_PATH")

    args = parser.parse_args()
    if args.date is None:
        args.date = datetime.now().strftime("%Y-%m-%d")

    if args.platform and args.platform_id:
        if args.platform == "CNVD":
            args.cnvd_id = args.platform_id
        elif args.platform == "CNNVD":
            args.cnnvd_id = args.platform_id

    update_summary(
        title=args.title,
        vendor=args.vendor,
        das_id=args.das_id,
        submitter=args.submitter,
        cnvd_id=args.cnvd_id,
        cnnvd_id=args.cnnvd_id,
        date=args.date,
        filepath=args.filepath,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
