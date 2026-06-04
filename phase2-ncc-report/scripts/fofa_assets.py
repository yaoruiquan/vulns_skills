#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""FOFA Web asset query helpers for NCC report preparation."""

from __future__ import annotations

import argparse
import base64
import json
import os
import re
import time
import urllib.parse
import urllib.request
from pathlib import Path
from typing import Any, Dict, Iterable

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm


SKILL_DIR = Path(__file__).resolve().parents[1]
ENV_FILE = SKILL_DIR / ".env"


def load_env() -> None:
    if not ENV_FILE.exists():
        return
    for raw_line in ENV_FILE.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        if key and key not in os.environ:
            os.environ[key] = value


load_env()


def env_bool(name: str, default: bool = False) -> bool:
    raw = os.environ.get(name)
    if raw is None:
        return default
    return raw.strip().lower() in {"1", "true", "yes", "y", "on", "是"}


def clean_product_name(value: str) -> str:
    text = re.sub(r"\s+", " ", value or "").strip()
    text = re.sub(r"(?i)\b(?:v|version)?\d+(?:\.\d+){1,5}(?:[-_][\w.]+)?\b", "", text)
    text = re.sub(r"(?i)\b(?:系统|平台|软件|项目|框架)?\s*(?:<=|>=|<|>)\s*[\w.:-]+", "", text)
    text = text.strip(" -_，,。:：")
    return text


def split_title_product(title: str) -> str:
    text = clean_product_name(title)
    text = re.sub(
        r"(?:存在|中的|的)?(?:SQL注入|XML实体注入|XXE|XSS|SSRF|CSRF|弱口令|文件上传|信息泄露|未授权访问|权限绕过|逻辑缺陷|文件包含|远程命令执行|命令执行|远程代码执行|任意代码执行|目录遍历|任意文件下载|任意文件读取|拒绝服务|反序列化|越界写入|越界读取|缓冲区溢出|二进制).*",
        "",
        text,
        flags=re.IGNORECASE,
    )
    return clean_product_name(text)


def product_variants(product: str) -> list[str]:
    base = clean_product_name(product)
    variants = [base] if base else []
    stripped = re.sub(r"(?:系统|平台|软件|项目|框架|组件|模块)$", "", base).strip(" -_，,。:：")
    if stripped and stripped not in variants:
        variants.append(stripped)
    spaced = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", stripped)
    if spaced and spaced not in variants:
        variants.append(spaced)
    return [item for item in variants if item]


def shell_quote(value: str) -> str:
    return '"' + value.replace("\\", "\\\\").replace('"', '\\"') + '"'


def fofa_query_candidates(product: str, vendor: str = "", title: str = "") -> list[str]:
    """Generate no-OR FOFA query candidates, most precise first."""
    product_names = product_variants(product) or product_variants(split_title_product(title))
    vendor_name = clean_product_name(vendor)
    candidates: list[str] = []

    for product_name in product_names:
        candidates.extend(
            [
                f'app={shell_quote(product_name)}',
                f'product={shell_quote(product_name)}',
                f'title={shell_quote(product_name)}',
                f'header={shell_quote(product_name)}',
                shell_quote(product_name),
                f'body={shell_quote(product_name)}',
            ]
        )
        if vendor_name and vendor_name not in product_name:
            candidates.append(f'(title={shell_quote(product_name)} && body={shell_quote(vendor_name)})')
            candidates.append(f'(body={shell_quote(product_name)} && body={shell_quote(vendor_name)})')
    if vendor_name and not product_names:
        candidates.append(f'body={shell_quote(vendor_name)}')

    deduped: list[str] = []
    seen = set()
    for query in candidates:
        compact = query.strip()
        if compact and "||" not in compact and " or " not in compact.lower() and compact not in seen:
            deduped.append(compact)
            seen.add(compact)
    return deduped or ['body="相关产品"']


def qbase64(query: str) -> str:
    return base64.b64encode(query.encode("utf-8")).decode("ascii")


def fofa_web_url(query: str) -> str:
    return "https://fofa.info/result?" + urllib.parse.urlencode({"qbase64": qbase64(query)})


def fofa_api_search(query: str) -> dict[str, Any]:
    """Optional paid API lookup. Disabled by default because Web quota is free."""
    if not env_bool("FOFA_API_ENABLED", False):
        return {"ok": False, "query": query, "reason": "FOFA_API_ENABLED=false，默认使用 Web 查询"}
    key = os.environ.get("FOFA_API_KEY", "").strip()
    if not key:
        return {"ok": False, "query": query, "reason": "FOFA_API_KEY 未配置"}
    try:
        timeout = float(os.environ.get("FOFA_TIMEOUT", "12") or "12")
    except ValueError:
        timeout = 12.0
    params = {
        "key": key,
        "qbase64": qbase64(query),
        "size": "1",
        "fields": "host,ip,port,title",
    }
    email = os.environ.get("FOFA_EMAIL", "").strip()
    if email:
        params["email"] = email
    url = "https://fofa.info/api/v1/search/all?" + urllib.parse.urlencode(params)
    request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 NCC report asset query"})
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            payload = json.loads(response.read().decode("utf-8", errors="ignore"))
    except Exception as exc:
        return {"ok": False, "query": query, "reason": str(exc)}
    if payload.get("error"):
        return {"ok": False, "query": query, "reason": payload.get("errmsg") or payload.get("error"), "payload": payload}
    return {
        "ok": True,
        "query": query,
        "size": int(payload.get("size") or 0),
        "payload": {key: value for key, value in payload.items() if key != "results"},
    }


def choose_fofa_query(candidates: Iterable[str]) -> dict[str, Any]:
    candidate_list = list(candidates)
    results = [fofa_api_search(query) for query in candidate_list]
    ok_results = [item for item in results if item.get("ok")]
    if ok_results:
        chosen = dict(next((item for item in results if item.get("ok") and int(item.get("size") or 0) > 0), ok_results[0]))
        chosen["attempts"] = results
        return chosen
    return {
        "ok": False,
        "query": candidate_list[0] if candidate_list else "",
        "size": "",
        "reason": "; ".join(dict.fromkeys(str(item.get("reason", "")) for item in results if item.get("reason"))) or "FOFA 查询失败",
        "attempts": results,
    }


def fofa_mcp_extract_script() -> str:
    """Return JS snippet to run on the FOFA result page via Chrome DevTools MCP."""
    return """() => {
  const text = document.body.innerText || "";
  const compact = text.replace(/\\s+/g, " ");
  const noResultPatterns = [
    /未找到相关结果/,
    /没有找到相关结果/,
    /没有找到匹配结果/,
    /暂无数据/,
    /No\\s+Results?/i,
    /No\\s+matching\\s+results?/i
  ];
  const noResult = noResultPatterns.some(pattern => pattern.test(compact));
  const patterns = [
    /(?:共|为您找到|查询到|匹配到|搜索到)[^0-9]{0,20}([0-9][0-9,]*)\\s*(?:条|个|项|资产|结果)/,
    /([0-9][0-9,]*)\\s*(?:条|个|项)?\\s*(?:匹配结果|搜索结果|相关资产|资产)/,
    /独立IP[^0-9]{0,20}([0-9][0-9,]*)/,
    /结果[^0-9]{0,20}([0-9][0-9,]*)/
  ];
  let assetCount = "";
  for (const pattern of patterns) {
    const match = compact.match(pattern);
    if (match) {
      assetCount = match[1].replace(/,/g, "");
      break;
    }
  }
  const title = document.title || "";
  const resultRows = document.querySelectorAll('table tbody tr, .el-table__row, [class*="result"] tbody tr').length;
  const hasResults = Boolean(assetCount && assetCount !== "0") && !noResult;
  return {
    ok: hasResults,
    has_results: hasResults,
    no_result: noResult,
    asset_count: assetCount,
    result_rows: resultRows,
    title,
    url: location.href,
    hint: hasResults ? "将 asset_count 写入 NCC Word，并保留截图" : "未检索到内容或未识别数量，不要保留截图",
    excerpt: compact.slice(0, 800)
  };
}"""


def screenshot_dir() -> Path:
    base = os.environ.get("FOFA_SCREENSHOT_DIR") or "/tmp/vulns-skills/phase2-ncc-report/fofa-screenshots"
    path = Path(base).expanduser()
    path.mkdir(parents=True, exist_ok=True)
    return path


def screenshot_match_key(value: str) -> str:
    text = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", value or "")
    text = text.lower()
    return re.sub(r"[^\w\u4e00-\u9fff]+", "", text)


def screenshot_hints(product: str = "", vendor: str = "", title: str = "", query: str = "") -> list[str]:
    values: list[str] = []
    values.extend(product_variants(product))
    if title:
        values.extend(product_variants(split_title_product(title)))
    if query:
        cleaned_query = re.sub(r"\b(?:app|product|title|header|body)\s*=", "", query)
        cleaned_query = cleaned_query.replace("&&", " ").strip("() ")
        values.extend(product_variants(cleaned_query))
    vendor_name = clean_product_name(vendor)
    if vendor_name and not values:
        values.append(vendor_name)
    deduped: list[str] = []
    seen = set()
    for value in values:
        key = screenshot_match_key(value)
        if key and key not in seen:
            deduped.append(value)
            seen.add(key)
    return deduped


def screenshot_files() -> list[Path]:
    files: list[Path] = []
    for suffix in ("*.png", "*.jpg", "*.jpeg", "*.webp"):
        files.extend(screenshot_dir().glob(suffix))
    return files


def recent_screenshot_files() -> list[Path]:
    files = screenshot_files()
    max_age = int(os.environ.get("FOFA_SCREENSHOT_MAX_AGE_SECONDS", "7200") or "7200")
    if max_age <= 0:
        return files
    return [path for path in files if time.time() - path.stat().st_mtime <= max_age]


def existing_screenshot_path(hints: Iterable[str] = (), allow_unmatched_latest: bool = False) -> str:
    explicit = os.environ.get("FOFA_SCREENSHOT_PATH", "").strip()
    if explicit and Path(explicit).expanduser().is_file():
        return str(Path(explicit).expanduser())
    if not env_bool("FOFA_USE_LATEST_SCREENSHOT", True):
        return ""
    files = screenshot_files()
    if not files:
        return ""
    hint_keys = [screenshot_match_key(hint) for hint in hints if screenshot_match_key(hint)]
    if hint_keys:
        matched = [
            path
            for path in files
            if any(key in screenshot_match_key(path.stem) or screenshot_match_key(path.stem) in key for key in hint_keys)
        ]
        if matched:
            return str(max(matched, key=lambda path: path.stat().st_mtime))
    if allow_unmatched_latest:
        recent_files = recent_screenshot_files()
        if recent_files:
            return str(max(recent_files, key=lambda path: path.stat().st_mtime))
    return ""


def build_asset_evidence(product: str, vendor: str = "", title: str = "") -> dict[str, Any]:
    candidates = fofa_query_candidates(product=product, vendor=vendor, title=title)
    if env_bool("FOFA_API_ENABLED", False):
        chosen = choose_fofa_query(candidates)
    else:
        chosen = {
            "ok": False,
            "query": candidates[0],
            "size": "",
            "reason": "默认使用 FOFA Web 查询免费额度；请用 Chrome DevTools MCP 打开 web_url 后提取数量",
        }
    query = str(chosen.get("query") or candidates[0])
    asset_count = str(chosen.get("size") if chosen.get("ok") else "")
    has_results = asset_has_results(asset_count)
    screenshot_path = existing_screenshot_path(
        screenshot_hints(product=product, vendor=vendor, title=title, query=query),
        allow_unmatched_latest=env_bool("FOFA_ALLOW_UNMATCHED_LATEST_SCREENSHOT", False),
    ) if has_results else ""
    return {
        "query": query,
        "qbase64": qbase64(query),
        "web_url": fofa_web_url(query),
        "asset_count": asset_count,
        "asset_has_results": has_results,
        "asset_valid": has_results,
        "api": chosen,
        "screenshot": {"ok": True, "path": screenshot_path} if screenshot_path else {},
        "mcp_extract_script": fofa_mcp_extract_script(),
        "candidates": candidates,
    }


def copy_run_format(source, target) -> None:
    target.bold = source.bold
    target.italic = source.italic
    target.underline = source.underline
    target.style = source.style
    target.font.name = source.font.name
    target.font.size = source.font.size
    target.font.color.rgb = source.font.color.rgb


def find_body_format_run(doc) -> Any:
    for index in (8, 12, 20, 22, 24):
        if index < len(doc.paragraphs):
            for run in doc.paragraphs[index].runs:
                if run.font.name or run.font.size or run.font.color.rgb:
                    return run
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.font.name or run.font.size or run.font.color.rgb:
                return run
    return None


def clear_paragraph_runs(paragraph) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._element)


def cleanup_unreferenced_images(doc) -> None:
    used_rel_ids = set()
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            for blip in run.element.xpath(".//a:blip"):
                rel_id = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                if rel_id:
                    used_rel_ids.add(rel_id)
    for rel_id, rel in list(doc.part.rels.items()):
        if rel.reltype == RT.IMAGE and rel_id not in used_rel_ids:
            del doc.part.rels[rel_id]


def set_paragraph_text(paragraph, text: str, format_run=None):
    clear_paragraph_runs(paragraph)
    run = paragraph.add_run(text)
    if format_run is not None:
        copy_run_format(format_run, run)
    return run


def parse_asset_count(asset_count: str) -> int | None:
    text = re.sub(r"\s+", " ", asset_count or "").strip()
    match = re.search(r"size\s*=\s*([0-9][0-9,]*)", text, flags=re.IGNORECASE)
    if match:
        return int(match.group(1).replace(",", ""))
    if re.fullmatch(r"[0-9][0-9,]*", text):
        return int(text.replace(",", ""))
    return None


def asset_has_results(asset_count: str) -> bool:
    count = parse_asset_count(asset_count)
    return count is not None and count > 0


def normalize_asset_count_text(asset_count: str) -> str:
    text = re.sub(r"\s+", " ", asset_count or "").strip()
    count = parse_asset_count(text)
    if count is not None and count > 0:
        return f"通过 FOFA 测绘查询，返回互联网资产数量 size={count}。"
    if count is not None:
        return f"未检索到互联网资产（FOFA 返回 size={count}）。"
    if not text:
        return "待通过 FOFA Web 查询后回填。"
    if text.startswith("通过 FOFA"):
        return text.rstrip("。") + "。"
    return f"通过 FOFA Web 测绘查询，返回互联网资产数量 {text}。"


def parse_optional_bool(value: str) -> bool | None:
    text = (value or "").strip().lower()
    if not text:
        return None
    if text in {"1", "true", "yes", "y", "on", "是"}:
        return True
    if text in {"0", "false", "no", "n", "off", "否"}:
        return False
    raise ValueError(f"无法识别布尔值: {value}")


def replace_asset_paragraphs(
    docx_path: Path,
    query: str,
    asset_count: str,
    screenshot_path: str = "",
    has_results: bool | None = None,
) -> None:
    doc = Document(str(docx_path))
    format_run = find_body_format_run(doc)
    screenshot = Path(screenshot_path).expanduser() if screenshot_path else Path()
    resolved_has_results = asset_has_results(asset_count) if has_results is None else has_results
    keep_screenshot = resolved_has_results and screenshot.is_file()
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("1）精准测绘语法"):
            set_paragraph_text(paragraph, f"1）精准测绘语法：{query}", format_run)
        elif text.startswith("2）互联网影响资产数量"):
            count_text = normalize_asset_count_text(asset_count)
            if not resolved_has_results:
                count_text = "未检索到互联网资产（FOFA 页面未返回与该检索语法匹配的有效结果）。"
            set_paragraph_text(paragraph, f"2）互联网影响资产数量：{count_text}", format_run)
        elif text.startswith("3）测绘平台截图"):
            set_paragraph_text(
                paragraph,
                "3）测绘平台截图：FOFA Web 查询结果截图如下。" if keep_screenshot else "3）测绘平台截图：无",
                format_run,
            )
            if keep_screenshot:
                run = paragraph.add_run()
                if format_run is not None:
                    copy_run_format(format_run, run)
                run.add_break()
                run.add_picture(str(screenshot), width=Cm(15))
    cleanup_unreferenced_images(doc)
    doc.save(str(docx_path))


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate FOFA asset evidence for NCC reports")
    subparsers = parser.add_subparsers(dest="command")

    generate = subparsers.add_parser("generate", help="Generate FOFA Web query URL and MCP extraction script")
    generate.add_argument("--product", default="")
    generate.add_argument("--vendor", default="")
    generate.add_argument("--title", default="")

    apply = subparsers.add_parser("apply", help="Write FOFA Web result back to an NCC Word document")
    apply.add_argument("--docx", required=True)
    apply.add_argument("--query", required=True)
    apply.add_argument("--asset-count", required=True)
    apply.add_argument("--screenshot-path", default="")
    apply.add_argument("--has-results", default="", help="MCP 页面提取结果：true/false。false 时不保留截图，即使传了截图路径。")

    parser.add_argument("--product", default="", help=argparse.SUPPRESS)
    parser.add_argument("--vendor", default="", help=argparse.SUPPRESS)
    parser.add_argument("--title", default="", help=argparse.SUPPRESS)
    args = parser.parse_args()
    if args.command == "apply":
        screenshot_path = args.screenshot_path or existing_screenshot_path(
            screenshot_hints(query=args.query, title=Path(args.docx).stem),
            allow_unmatched_latest=True,
        )
        has_results = parse_optional_bool(args.has_results)
        replace_asset_paragraphs(Path(args.docx).expanduser(), args.query, args.asset_count, screenshot_path, has_results)
        print(
            json.dumps(
                {"ok": True, "docx": args.docx, "screenshot_path": screenshot_path, "has_results": has_results},
                ensure_ascii=False,
            )
        )
        return 0
    print(json.dumps(build_asset_evidence(args.product, args.vendor, args.title), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
