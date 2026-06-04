#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成 NCC 上报前材料目录和通用型漏洞报告 Word。"""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
from datetime import datetime
from pathlib import Path
from typing import Dict, Optional

from docx import Document
from docx.shared import Cm

from extract_vuln_data import (
    DEFAULT_DATA_DIR,
    clean_text_for_textarea,
    classify_detail,
    detect_material_source,
    extract_das_id,
    extract_fields_from_docx,
    find_material_dir,
    find_preferred_docx,
    first_value,
    infer_vendor,
    is_valid_docx,
    normalize_yes_no,
    normalize_product_version,
    normalize_text,
    resolve_input,
    sanitize_filename,
)
from fofa_assets import build_asset_evidence
from web_enrichment import build_security_guidance


DEFAULT_TEMPLATE_PATH = Path(
    os.environ.get(
        "NCC_REPORT_TEMPLATE_PATH",
        "/Users/yao/Documents/网安- AI应用开发/监管上报/NCC通用型漏洞报告模板.docx",
    )
).expanduser()
PLATFORM_PREFIX_RE = re.compile(r"^(?:CNVD|CNNVD|NCC)-", re.IGNORECASE)
PLACEHOLDER_HEADINGS = {
    "漏洞基本信息",
    "提交人员基本信息",
    "漏洞验证",
    "漏洞验证过程",
    "漏洞分析",
    "漏洞描述",
    "漏洞危害",
    "修复方案",
    "临时解决方案",
    "正式解决方案",
    "复现注意事项",
}
ATTACHMENT_NOTE_RE = re.compile(r"(详见|见)附件[^。；\n]*(?:[。；])?")


def resolve_bound_day_original(args: argparse.Namespace) -> str:
    """是否0Day漏洞和是否原创漏洞在 NCC 流程中绑定为同一个是/否值。"""
    values = []
    for value in (args.is_0day_original, args.is_0day, args.is_original):
        normalized = normalize_yes_no(value, default="")
        if normalized:
            values.append(normalized)
    if not values:
        raise SystemExit(
            "生成 NCC 材料前必须先确认是否0Day/原创绑定值，请传入 --is-0day-original 是 或 --is-0day-original 否"
        )
    if any(value not in {"是", "否"} for value in values) or len(set(values)) != 1:
        raise SystemExit("是否0Day漏洞和是否原创漏洞必须绑定一致：是都为是，否都为否")
    return values[0]


def valid_value(value: str) -> bool:
    """判断字段是否可以直接写入 NCC 报告。"""
    text = normalize_text(value)
    compact = re.sub(r"\s+", "", text).strip(":：")
    return bool(text) and text not in {"无", "暂无", "N/A", "n/a", "none", "None", "-", "--"} and compact not in PLACEHOLDER_HEADINGS


def merge_fields(primary: Dict[str, str], fallback: Dict[str, str]) -> Dict[str, str]:
    """CNVD 为主，CNNVD 作为缺失字段补充。"""
    merged = dict(fallback)
    for key, value in primary.items():
        if valid_value(value) or key not in merged:
            merged[key] = value
    return merged


def find_platform_docx(das_root: Path, platform: str) -> tuple[Optional[Path], Dict[str, str]]:
    """查找指定平台 Word 并提取字段。"""
    material_dir = find_material_dir(das_root, platform)
    if not material_dir:
        return None, {}
    docx_path = find_preferred_docx(material_dir)
    if not docx_path:
        return None, {}
    return docx_path, extract_fields_from_docx(docx_path)


def resolve_das_root(target: Path, prefer_source: str) -> tuple[Path, Path, Path]:
    """兼容 DAS 目录、平台目录和 docx 输入，返回 DAS 根、主材料目录和主 Word。"""
    args = argparse.Namespace(
        das_id="" if target.exists() else str(target),
        data_dir=DEFAULT_DATA_DIR,
        input_path=str(target) if target.exists() else "",
        docx_path=str(target) if target.is_file() else "",
        prefer_source=prefer_source,
        platform="NCC",
    )
    material_dir, docx_path, error = resolve_input(args)
    if error:
        raise SystemExit(error)
    assert material_dir is not None
    assert docx_path is not None
    source = detect_material_source(material_dir)
    das_root = material_dir.parent if source in {"CNVD", "CNNVD", "NCC"} else material_dir
    return das_root, material_dir, docx_path


def material_name(title: str, source_dir: Path) -> str:
    """生成 NCC 材料目录/Word 文件名。"""
    name = source_dir.name
    if detect_material_source(source_dir) in {"CNVD", "CNNVD", "NCC"}:
        name = PLATFORM_PREFIX_RE.sub("", name)
    if not name:
        name = title
    return sanitize_filename(name, sanitize_filename(title, "NCC漏洞材料"))


def copy_run_format(source, target) -> None:
    """复制模板 run 的基础字符格式。"""
    target.bold = source.bold
    target.italic = source.italic
    target.underline = source.underline
    target.style = source.style
    target.font.name = source.font.name
    target.font.size = source.font.size
    target.font.color.rgb = source.font.color.rgb


def replace_paragraph_body(paragraph, text: str, format_run=None) -> None:
    """替换正文段落内容，保留原段落和首个 run 的格式。"""
    if not paragraph.runs:
        paragraph.add_run(text)
        if format_run is not None:
            copy_run_format(format_run, paragraph.runs[0])
        return
    paragraph.runs[0].text = text
    if format_run is not None:
        copy_run_format(format_run, paragraph.runs[0])
    for run in paragraph.runs[1:]:
        run.text = ""


def replace_after_kept_runs(paragraph, keep_count: int, text: str, format_run=None) -> None:
    """保留字段名前缀 run，只替换其后的模板要求文字。"""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        if format_run is not None:
            copy_run_format(format_run, paragraph.runs[0])
        return
    keep_count = min(keep_count, len(runs))
    value_index = next((index for index in range(keep_count, len(runs)) if not runs[index].bold), keep_count)
    value_run = runs[value_index] if value_index < len(runs) else paragraph.add_run()
    for run in runs[keep_count:]:
        run.text = ""
    value_run.text = text
    if format_run is not None:
        copy_run_format(format_run, value_run)


def append_value_run(paragraph, text: str, template_run_index: int = -1, format_run=None) -> None:
    """在原字段标题后追加值，并复用模板字符格式。"""
    if not paragraph.runs:
        paragraph.add_run(text)
        if format_run is not None:
            copy_run_format(format_run, paragraph.runs[0])
        return
    source = format_run or paragraph.runs[template_run_index]
    run = paragraph.add_run(text)
    copy_run_format(source, run)


def clear_paragraph(paragraph) -> None:
    """清空模板示例文字但保留段落位置和格式。"""
    for run in paragraph.runs:
        run.text = ""


def paragraph_value(value: str, fallback: str = "无") -> str:
    """输出适合 Word 段落的值。"""
    text = clean_text_for_textarea(value)
    return text if valid_value(text) else fallback


def first_clean_value(fields: Dict[str, str], *keys: str, max_chars: int = 0) -> str:
    """按候选字段读取清洗后仍非空的值。"""
    for key in keys:
        value = clean_text_for_textarea(fields.get(key, ""), max_chars=max_chars)
        if valid_value(value):
            return value
    return ""


def dedupe_product_version(product: str, version: str) -> str:
    """组合产品和版本，避免 xl2tpd 1.3.20 1.3.20 这类重复。"""
    product = normalize_text(product)
    version = normalize_text(version)
    if not product:
        return version
    if not version or version in product:
        return product
    return f"{product} {version}".strip()


def product_without_version(product: str, version: str) -> str:
    """产品介绍中尽量使用纯产品名。"""
    product = normalize_text(product)
    version = normalize_text(version)
    if version and product.endswith(version):
        return product[: -len(version)].strip()
    return product


def product_intro(unit_name: str, product: str) -> str:
    """生成简短产品介绍，避免模板示例残留。"""
    product_name = product or unit_name or "相关产品"
    vendor = unit_name or "相关厂商"
    return f"{product_name} 是 {vendor} 相关的软件或系统组件，具体产品形态及部署环境以附件证明材料为准。"


def build_asset_values(product: str, unit_name: str, title: str) -> Dict[str, str]:
    """生成 FOFA 测绘语法、资产数量和截图路径。"""
    evidence = build_asset_evidence(product=product, vendor=unit_name, title=title)
    query = normalize_text(str(evidence.get("query") or ""))
    count = normalize_text(str(evidence.get("asset_count") or ""))
    asset_has_results = bool(evidence.get("asset_has_results") or evidence.get("asset_valid"))
    screenshot = evidence.get("screenshot") if isinstance(evidence.get("screenshot"), dict) else {}
    screenshot_path = normalize_text(str(screenshot.get("path") or "")) if screenshot.get("ok") else ""
    if count and count != "无" and asset_has_results:
        count_text = f"通过 FOFA 测绘查询，返回互联网资产数量 size={count}。"
    elif count and count != "无":
        count_text = f"未检索到互联网资产（FOFA 返回 size={count}）。"
    else:
        count_text = "待通过 FOFA Web 查询后回填。"
    return {
        "asset_query": query or "无",
        "asset_count": count_text,
        "asset_screenshot": "FOFA 查询结果截图如下。" if screenshot_path and asset_has_results else "无" if count else "待通过 FOFA Web 查询后回填。",
        "asset_screenshot_path": screenshot_path if asset_has_results else "",
        "asset_evidence": evidence,
    }


def impact_text(fields: Dict[str, str], description: str, detail_category: str) -> str:
    """提炼漏洞危害，缺失时给出保守描述。"""
    direct = first_value(fields, "漏洞危害")
    if valid_value(direct):
        return clean_text_for_textarea(direct, max_chars=700)
    if valid_value(description):
        summary = clean_text_for_textarea(description, max_chars=260)
        summary = ATTACHMENT_NOTE_RE.sub("", summary).strip()
        return f"攻击者利用该{detail_category or '漏洞'}可能影响相关系统的安全性。{summary}具体影响以附件中的验证材料为准。"
    return "无"


def solution_text(fields: Dict[str, str]) -> str:
    """生成修复方案文本。"""
    raw = first_value(fields, "修复方案", "临时解决方案", "正式解决方案")
    if valid_value(raw):
        return clean_text_for_textarea(raw, max_chars=700)
    return "建议升级至官方修复版本；如暂无补丁，临时限制相关接口访问并加强输入校验。"


def needs_guidance_update(docx_path: Path) -> bool:
    """已有 NCC Word 缺少危害/修复方案时允许准备阶段补齐。"""
    if not docx_path.exists():
        return True
    fields = extract_fields_from_docx(docx_path)
    impact = first_value(fields, "漏洞危害", "危害说明", "影响说明")
    solution = first_value(fields, "修复方案", "临时解决方案", "正式解决方案", "修复建议", "解决方案")
    return not valid_value(impact) or not valid_value(solution)


def needs_asset_update(docx_path: Path) -> bool:
    """已有 NCC Word 资产证明还是模板/空值时允许准备阶段刷新。"""
    if not docx_path.exists():
        return True
    doc = Document(str(docx_path))
    checks = {
        "query": False,
        "count": False,
        "screenshot": False,
    }
    screenshot_has_image = False
    for paragraph in doc.paragraphs:
        text = normalize_text(paragraph.text)
        if text.startswith("1）精准测绘语法"):
            checks["query"] = valid_value(text.split("：", 1)[-1] if "：" in text else text)
        elif text.startswith("2）互联网影响资产数量"):
            value = text.split("：", 1)[-1] if "：" in text else text
            checks["count"] = valid_value(value) and "待通过" not in value and "未获取" not in value
        elif text.startswith("3）测绘平台截图"):
            value = text.split("：", 1)[-1] if "：" in text else text
            checks["screenshot"] = valid_value(value) and value != "无"
            screenshot_has_image = any(run.element.xpath(".//pic:pic") for run in paragraph.runs)
    return not checks["query"] or not checks["count"] or not checks["screenshot"] or not screenshot_has_image


def split_verification(fields: Dict[str, str]) -> tuple[str, str]:
    """把验证过程拆成环境搭建和触发操作，无法拆分时保守填充。"""
    text = first_clean_value(fields, "漏洞验证过程", "验证过程", "漏洞验证", "漏洞分析", max_chars=1600)
    if not valid_value(text):
        return "见附件", "见附件"
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    env_lines = [line for line in lines if re.search(r"环境|部署|安装|编译|docker|源码|版本|复用共享环境", line, re.IGNORECASE)]
    if env_lines:
        env = "\n".join(env_lines[:6])
        return env, text
    return "见附件", text


def build_report_values(fields: Dict[str, str], source_dir: Path) -> Dict[str, str]:
    """从 CNVD/CNNVD 字段构造 NCC 模板写入值。"""
    title = clean_text_for_textarea(first_value(fields, "漏洞名称"), max_chars=180)
    description = clean_text_for_textarea(first_value(fields, "漏洞描述", "漏洞简介"), max_chars=1000)
    product, version = normalize_product_version(fields, title)
    unit_name = infer_vendor(fields, title, product)
    detail_category = classify_detail(fields, title, description)
    env_text, trigger_operation = split_verification(fields)
    trigger_condition = first_value(fields, "漏洞触发条件", "触发条件")
    report_title = title or material_name("", source_dir)
    clean_product = product_without_version(product, version)
    product_version = dedupe_product_version(product, version)
    poc = first_clean_value(fields, "漏洞验证过程", "验证过程", "漏洞验证", "漏洞分析", max_chars=1600)
    guidance = build_security_guidance(
        fields=fields,
        title=report_title,
        product=product,
        vendor=unit_name,
        category=detail_category,
        description=description,
    )
    asset_values = build_asset_values(clean_product or product, unit_name, report_title)
    return {
        "title": report_title,
        "discovery_date": datetime.now().strftime("%Y年%m月%d日"),
        "detail_category": detail_category or paragraph_value(first_value(fields, "漏洞类型"), "其他"),
        "description_intro": product_intro(unit_name, clean_product),
        "description": paragraph_value(description),
        "impact": str(guidance.get("impact") or impact_text(fields, description, detail_category)),
        "unit_name": paragraph_value(unit_name),
        "product_version": paragraph_value(product_version, "暂未明确"),
        "asset_query": asset_values["asset_query"],
        "asset_count": asset_values["asset_count"],
        "asset_screenshot": asset_values["asset_screenshot"],
        "asset_screenshot_path": asset_values["asset_screenshot_path"],
        "asset_evidence": asset_values["asset_evidence"],
        "poc": paragraph_value(poc, "见附件"),
        "trigger_condition": paragraph_value(trigger_condition, "见附件"),
        "environment": paragraph_value(env_text, "见附件"),
        "trigger_operation": paragraph_value(trigger_operation, "见附件"),
        "reproduce_note": "无",
        "blackbox_cases": "无",
        "affected_targets": "无",
        "solution": str(guidance.get("solution") or solution_text(fields)),
        "security_guidance": guidance,
        "remark": "本材料由 CNVD/CNNVD 原始材料整理生成，证明材料见附件。",
    }


def fill_template(template_path: Path, output_path: Path, values: Dict[str, str]) -> None:
    """在 NCC 模板原格式上替换占位说明和示例内容。"""
    doc = Document(str(template_path))
    paragraphs = doc.paragraphs
    if len(paragraphs) < 63:
        raise SystemExit(f"NCC Word 模板段落数量异常，无法按模板填充: {template_path}")
    body_run = paragraphs[8].runs[0] if paragraphs[8].runs else None

    replace_after_kept_runs(paragraphs[0], 1, values["title"], format_run=body_run)
    append_value_run(paragraphs[2], values["discovery_date"], format_run=body_run)
    append_value_run(paragraphs[4], values["detail_category"], format_run=body_run)
    clear_paragraph(paragraphs[5])

    replace_after_kept_runs(paragraphs[7], 1, "")
    replace_paragraph_body(paragraphs[8], values["description_intro"], format_run=body_run)
    replace_paragraph_body(paragraphs[9], values["description"], format_run=body_run)

    replace_after_kept_runs(paragraphs[11], 2, "")
    replace_paragraph_body(paragraphs[12], values["impact"], format_run=body_run)

    replace_after_kept_runs(paragraphs[14], 1, values["unit_name"], format_run=body_run)
    paragraphs[16].runs[2].text = ""
    paragraphs[16].runs[3].text = ""
    paragraphs[16].runs[4].text = ""
    paragraphs[16].runs[5].text = ":"
    append_value_run(paragraphs[16], values["product_version"], format_run=body_run)
    clear_paragraph(paragraphs[17])

    replace_paragraph_body(paragraphs[20], f"1）精准测绘语法：{values['asset_query']}", format_run=body_run)
    replace_paragraph_body(paragraphs[22], f"2）互联网影响资产数量：{values['asset_count']}", format_run=body_run)
    replace_paragraph_body(paragraphs[24], f"3）测绘平台截图：{values['asset_screenshot']}", format_run=body_run)
    screenshot_path = Path(str(values.get("asset_screenshot_path") or "")).expanduser()
    if screenshot_path.is_file():
        run = paragraphs[24].add_run()
        run.add_break()
        run.add_picture(str(screenshot_path), width=Cm(15))

    replace_paragraph_body(paragraphs[32], values["poc"], format_run=body_run)
    for index in (33, 34, 35, 36, 37):
        clear_paragraph(paragraphs[index])

    replace_paragraph_body(paragraphs[40], values["trigger_condition"], format_run=body_run)
    replace_paragraph_body(paragraphs[45], values["environment"], format_run=body_run)
    replace_paragraph_body(paragraphs[48], values["trigger_operation"], format_run=body_run)
    replace_paragraph_body(paragraphs[51], values["reproduce_note"], format_run=body_run)

    replace_paragraph_body(paragraphs[54], f"具体案例证明（3个）：{values['blackbox_cases']}", format_run=body_run)
    replace_paragraph_body(paragraphs[56], f"其他受影响目标（10个以上）：{values['affected_targets']}", format_run=body_run)
    clear_paragraph(paragraphs[57])

    replace_after_kept_runs(paragraphs[59], 2, "")
    replace_paragraph_body(paragraphs[60], values["solution"], format_run=body_run)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))


def copy_supporting_files(source_dir: Path, target_dir: Path) -> list[str]:
    """复制非 Word 附件和子目录，保持材料结构。"""
    copied: list[str] = []
    for source in sorted(source_dir.rglob("*")):
        if source.name.startswith(".") or source.name.startswith("~$"):
            continue
        if source.is_file() and source.suffix.lower() in {".docx", ".doc"}:
            continue
        relative = source.relative_to(source_dir)
        target = target_dir / relative
        if source.is_dir():
            target.mkdir(parents=True, exist_ok=True)
        elif source.is_file():
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, target)
        else:
            continue
        copied.append(str(target))
    return copied


def prepare_ncc_material(
    target: str,
    prefer_source: str = "CNVD",
    template_path: Path = DEFAULT_TEMPLATE_PATH,
    force: bool = False,
) -> Dict[str, object]:
    """生成 NCC 材料目录，返回路径和来源信息。"""
    target_path = Path(target).expanduser()
    das_root, source_dir, source_docx = resolve_das_root(target_path, prefer_source)
    if detect_material_source(source_dir) == "NCC" and source_docx.exists() and not force:
        updated_existing = False
        refresh_guidance = needs_guidance_update(source_docx)
        refresh_assets = needs_asset_update(source_docx)
        if template_path.exists() and (refresh_guidance or refresh_assets):
            values = build_report_values(extract_fields_from_docx(source_docx), source_dir)
            fill_template(template_path, source_docx, values)
            updated_existing = True
            reason = "already NCC material; refreshed"
            if refresh_guidance and refresh_assets:
                reason = "already NCC material; guidance and FOFA asset fields refreshed"
            elif refresh_guidance:
                reason = "already NCC material; guidance fields refreshed"
            elif refresh_assets:
                reason = "already NCC material; FOFA asset fields refreshed"
        else:
            reason = "already NCC material"
        return {
            "ok": True,
            "created": False,
            "updated_existing": updated_existing,
            "reason": reason,
            "das_root": str(das_root),
            "source_dir": str(source_dir),
            "source_docx": str(source_docx),
            "ncc_dir": str(source_dir),
            "ncc_docx": str(source_docx),
        }

    if not template_path.exists():
        raise SystemExit(f"NCC Word 模板不存在: {template_path}")

    cnvd_docx, cnvd_fields = find_platform_docx(das_root, "CNVD")
    cnnvd_docx, cnnvd_fields = find_platform_docx(das_root, "CNNVD")
    primary_fields = cnvd_fields or extract_fields_from_docx(source_docx)
    fields = merge_fields(primary_fields, cnnvd_fields)
    values = build_report_values(fields, source_dir)
    folder_stem = material_name(values["title"], source_dir)
    ncc_dir = das_root / f"NCC-{folder_stem}"
    ncc_docx = ncc_dir / f"NCC-{folder_stem}.docx"
    if ncc_dir.exists() and not force:
        existing_docx = find_preferred_docx(ncc_dir)
        if existing_docx:
            updated_existing = False
            refresh_guidance = needs_guidance_update(existing_docx)
            refresh_assets = needs_asset_update(existing_docx)
            if refresh_guidance or refresh_assets:
                fill_template(template_path, existing_docx, values)
                updated_existing = True
                reason = "NCC material exists; refreshed"
                if refresh_guidance and refresh_assets:
                    reason = "NCC material exists; guidance and FOFA asset fields refreshed"
                elif refresh_guidance:
                    reason = "NCC material exists; guidance fields refreshed"
                elif refresh_assets:
                    reason = "NCC material exists; FOFA asset fields refreshed"
            else:
                reason = "NCC material exists"
            return {
                "ok": True,
                "created": False,
                "updated_existing": updated_existing,
                "reason": reason,
                "das_root": str(das_root),
                "source_dir": str(source_dir),
                "source_docx": str(source_docx),
                "cnvd_docx": str(cnvd_docx or ""),
                "cnnvd_docx": str(cnnvd_docx or ""),
                "ncc_dir": str(ncc_dir),
                "ncc_docx": str(existing_docx),
                "values": values if updated_existing else {},
            }
        raise SystemExit(f"NCC 材料目录已存在但缺少 docx，请检查或使用 --force: {ncc_dir}")

    if ncc_dir.exists() and force:
        shutil.rmtree(ncc_dir)
    ncc_dir.mkdir(parents=True, exist_ok=True)
    copied = copy_supporting_files(source_dir, ncc_dir)
    fill_template(template_path, ncc_docx, values)
    return {
        "ok": True,
        "created": True,
        "das_root": str(das_root),
        "source_dir": str(source_dir),
        "source_docx": str(source_docx),
        "cnvd_docx": str(cnvd_docx or ""),
        "cnnvd_docx": str(cnnvd_docx or ""),
        "ncc_dir": str(ncc_dir),
        "ncc_docx": str(ncc_docx),
        "copied_files": copied,
        "values": values,
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="生成 NCC 上报前材料目录和通用型漏洞报告 Word")
    parser.add_argument("target", help="DAS-ID、DAS 目录、CNVD/CNNVD/NCC 目录或 docx 路径")
    parser.add_argument("--is-0day-original", default="", help="是否0Day漏洞/是否原创漏洞绑定值：是/否/yes/no/true/false/1/0；生成材料前必须明确传入")
    parser.add_argument("--is-0day", default="", help="兼容旧参数；必须与 --is-original 和 --is-0day-original 一致")
    parser.add_argument("--is-original", default="", help="兼容旧参数；必须与 --is-0day 和 --is-0day-original 一致")
    parser.add_argument("--prefer-source", choices=["CNVD", "CNNVD", "NCC"], default="CNVD", help="优先读取的平台材料")
    parser.add_argument("--template", default=str(DEFAULT_TEMPLATE_PATH), help="NCC 通用型漏洞报告模板 docx")
    parser.add_argument("--force", action="store_true", help="覆盖重建已有 NCC 材料目录")
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    resolve_bound_day_original(args)
    result = prepare_ncc_material(
        args.target,
        prefer_source=args.prefer_source,
        template_path=Path(args.template).expanduser(),
        force=args.force,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
